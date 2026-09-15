"""Create a demo Word template that uses every supported tag.

Run:
    python make_sample_template.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

OUTPUT = Path(__file__).with_name("templates") / "sample_template.docx"


def main() -> None:
    OUTPUT.parent.mkdir(exist_ok=True)

    doc = Document()
    doc.add_heading("Užsakymas Nr. {{ order_number }}", level=1)

    p = doc.add_paragraph()
    p.add_run("Data: ").bold = True
    p.add_run("{{ date }}")

    p = doc.add_paragraph()
    p.add_run("Užsakymo numeris: ").bold = True
    p.add_run("{{ order_number }}")

    p = doc.add_paragraph()
    p.add_run("Suma: ").bold = True
    p.add_run("{{ order_price_eur }}")

    p = doc.add_paragraph()
    p.add_run("Suma žodžiais: ").bold = True
    p.add_run("{{ order_price_words }}")

    doc.add_paragraph(
        "Šis dokumentas sugeneruotas automatiškai iš „Order document filler“."
    )

    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
